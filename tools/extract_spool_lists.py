# -*- coding: utf-8 -*-
"""
Traekker SPOOL-arkets vaerdilister og opslagstabeller ud til seed-CSV.

    python3 tools/extract_spool_lists.py <sti til FL_indberetninger...xlsm>

        -> sharepoint/seed/MD_FLValueList.csv   dropdown-vaerdier
        -> sharepoint/seed/MD_FLKey.csv         klassebestemmelsens opslag

    python3 tools/extract_spool_lists.py <xlsm> --write-js
        -> html/lookups.generated.js            FL_LOOKUPS (R2-076)

    python3 tools/extract_spool_lists.py <xlsm> --check-js
        sammenligner arket med den committede html/lookups.generated.js
        tabel for tabel; exit 1 ved forskel. Intet skrives.

FL_LOOKUPS FRA ARKET (R2-076)
------------------------------
lookups.generated.js er SPOOL-arkets navngivne tabeller (alle ark, i
arkets raekkefoelge): {"sheet", "headers", "rows"}, alle vaerdier som
tekst. Filen blev foer skrevet af Windows PowerShell 5.1 (ConvertTo-Json)
og er her efterlignet byte for byte: BOM, CRLF, PowerShells indrykning og
JavaScriptSerializers escaping (& < > som \\u0026 \\u003c \\u003e).
format_lookups_js(json.loads(<filen>)) giver filen igen
(tests/test_lookups_generator.py), saa en forskel efter --write-js
skyldes arket, ikke formatet.

Kaeden: arket -> --write-js -> lookups.generated.js -> node
tools/fl/harness.js plan (appens regler) og seed (MD_FLKey.csv), og
gen_kks_seed.py --check / harness.js test som tjek af, at de er i trit
(docs/31 PX5/AQ8).

HVORFOR TO LISTER OG IKKE EEN
------------------------------
De to slags data bruges forskellige steder. Vaerdilisterne fylder en
dropdown og valideres mod, naar brugeren har skrevet noget.
Opslagstabellerne laeses derimod ud fra selve KKS-koden, foer brugeren har
udfyldt noget som helst, og afgoer hvilken klasse raekken faar - og dermed
hvilke felter der overhovedet skal vises.

SCEq ER IKKE EN LISTE, MEN SEKSTEN
-----------------------------------
SCEq-tabellen har en kolonne pr. klasse, hver med sine egne gyldige
vaerdier. VBA'en slaar op med sceTable.ListColumns(Class). Den foldes ud
her, saa hver klasse faar sine egne raekker under ListName "SCEq:<klasse>"
- ellers ville en ELF kunne vaelge en MKP-vaerdi.
"""

import csv
import io
import os
import sys
import warnings

warnings.filterwarnings("ignore")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEED = os.path.join(ROOT, "sharepoint", "seed")

# Vaerdilister paa arket "List data". Navnet i arket -> navnet i appen.
# Table20 og Table10 hedder ikke noget i Excel; her faar de et navn, der
# siger hvad de er.
VALUE_TABLES = {
    "TypekredsTabel": "Typekreds",
    "TestMethod": "TestMethod",
    "FireClassification": "FireClassification",
    "Table20": "FireSealingType",
    "Design_pressure_uom": "Design_pressure_uom",
    "Design_Temp_uom": "Design_Temp_uom",
    "Design_flow_uom": "Design_flow_uom",
    "Plant": "Plant",
    "EQ_Cat": "EquipmentCategory",
}

# Opslagstabeller paa arket "DictionaryTable".
# navn -> (KeyType, kolonne med noeglen, kolonne med klassen/vaerdien)
KEY_TABLES = {
    "ClassDeterminationComponentKey": ("Component", 0, 1),
    "ClassDeterminationAggregateKey": ("Aggregate", 0, 1),
    "FunctionKeyDict": ("Function", 0, None),
    "BR18_Keys": ("BR18", 0, 1),
}


def write_seed(path, header, rows):
    """Samme format som de oevrige seed-filer: komma, BOM, alt citeret."""
    with io.open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, quoting=csv.QUOTE_ALL)
        w.writerow(header)
        w.writerows(rows)


def table_rows(wb, sheet, name):
    """En navngiven Excel-tabel -> liste af raekker, uden overskriften."""
    ws = wb[sheet]
    ref = ws.tables[name].ref
    rows = [[c.value for c in r] for r in ws[ref]]
    return rows[0], rows[1:]


def _clean(v):
    if v is None:
        return ""
    return str(v).strip()


# ---------------------------------------------------------------------------
# html/lookups.generated.js (R2-076)
# ---------------------------------------------------------------------------
JS_OUT = os.path.join(ROOT, "html", "lookups.generated.js")
JS_PREFIX = "window.FL_LOOKUPS = "


def _cell_text(v):
    """Cellens vaerdi som tekst, som PowerShell skrev den: tom celle -> "",
    et helt tal uden ",0"."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def lookups_from_workbook(wb):
    """{tabelnavn: {"sheet", "headers", "rows"}} i arkets raekkefoelge.
    Helt tomme raekker (fx en tabels tomme indtastningsraekke) udelades."""
    out = {}
    for ws in wb.worksheets:
        for name in ws.tables:
            cells = [[_cell_text(c.value) for c in r] for r in ws[ws.tables[name].ref]]
            out[name] = {"sheet": ws.title, "headers": cells[0],
                         "rows": [r for r in cells[1:] if any(r)]}
    return out


def _js_string(s):
    """En JSON-streng med JavaScriptSerializers escaping."""
    out = ['"']
    for ch in s:
        o = ord(ch)
        if ch == '"':
            out.append('\\"')
        elif ch == "\\":
            out.append("\\\\")
        elif ch in "<>&'" or o < 0x20:
            out.append({"\b": "\\b", "\f": "\\f", "\n": "\\n", "\r": "\\r",
                        "\t": "\\t"}.get(ch) or "\\u%04x" % o)
        else:
            out.append(ch)
    out.append('"')
    return "".join(out)


def _ps_json(v, col):
    """PowerShell 5.1's ConvertTo-Json-indrykning: et objekts/listes indhold
    staar 4 tegn inde i forhold til kolonnen, hvor '{' / '[' staar, og den
    lukkende parentes staar i samme kolonne som den aabnende. En tom liste
    er '[', en tom linje og ']'."""
    pad = " " * (col + 4)
    if isinstance(v, dict):
        parts = []
        for k, x in v.items():
            head = _js_string(k) + ":  "
            parts.append(pad + head + _ps_json(x, col + 4 + len(head)))
        return "{\r\n" + ",\r\n".join(parts) + "\r\n" + " " * col + "}"
    if isinstance(v, list):
        if not v:
            return "[\r\n\r\n" + " " * col + "]"
        return ("[\r\n" + ",\r\n".join(pad + _ps_json(x, col + 4) for x in v)
                + "\r\n" + " " * col + "]")
    return _js_string(str(v))


def format_lookups_js(lookups, eol="\r\n"):
    """Hele lookups.generated.js som bytes (UTF-8 med BOM). PowerShell skrev
    CRLF; .gitattributes (eol=lf) kan give LF i en checkout - eol foelger
    derfor den fil, der ligger der (js_eol)."""
    text = JS_PREFIX + _ps_json(lookups, 0) + ";\r\n"
    if eol != "\r\n":
        text = text.replace("\r\n", eol)
    return b"\xef\xbb\xbf" + text.encode("utf-8")


def js_eol(path=JS_OUT):
    """Linjeskiftet i den eksisterende fil (CRLF, hvis den ikke findes)."""
    try:
        with open(path, "rb") as f:
            head = f.read(4096)
    except OSError:
        return "\r\n"
    return "\r\n" if b"\r\n" in head or b"\n" not in head else "\n"


def read_lookups_js(path=JS_OUT):
    import json
    with open(path, encoding="utf-8-sig") as f:
        t = f.read()
    return json.loads(t[t.index("{"):t.rindex("}") + 1])


def check_js(lookups, path=JS_OUT):
    """Forskellene mellem arket og den committede fil, tabel for tabel."""
    old = read_lookups_js(path)
    diffs = []
    for name in list(old) + [n for n in lookups if n not in old]:
        if name not in lookups:
            diffs.append("%s: findes i lookups.generated.js, ikke i arket" % name)
        elif name not in old:
            diffs.append("%s: findes i arket, ikke i lookups.generated.js" % name)
        elif lookups[name] != old[name]:
            a, b = old[name], lookups[name]
            what = [k for k in ("sheet", "headers") if a.get(k) != b.get(k)]
            if a["rows"] != b["rows"]:
                what.append("rows (%d i filen, %d i arket)" % (len(a["rows"]), len(b["rows"])))
            diffs.append("%s: %s" % (name, ", ".join(what)))
    if list(old) != list(lookups) and not diffs:
        diffs.append("tabellernes raekkefoelge er en anden")
    return diffs


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if len(args) != 1 or flags - {"--write-js", "--check-js"}:
        sys.exit("brug: extract_spool_lists.py <sti til .xlsm> [--write-js | --check-js]")
    path = args[0]
    if not os.path.exists(path):
        sys.exit("findes ikke: %s" % path)

    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    if "--check-js" in flags:
        diffs = check_js(lookups_from_workbook(wb))
        for d in diffs:
            print("  " + d)
        print("lookups.generated.js er i trit med arket." if not diffs else
              "%d tabel(ler) afviger fra arket." % len(diffs))
        sys.exit(1 if diffs else 0)
    if "--write-js" in flags:
        data = format_lookups_js(lookups_from_workbook(wb), js_eol())
        with open(JS_OUT, "wb") as f:
            f.write(data)
        print("Skrev %s (%d bytes). Koer nu: node tools/fl/harness.js plan "
              "&& node tools/fl/harness.js seed" % (os.path.relpath(JS_OUT, ROOT), len(data)))
        return

    # ---- vaerdilister ----------------------------------------------------
    values = []
    for xl_name, list_name in sorted(VALUE_TABLES.items()):
        head, rows = table_rows(wb, "List data", xl_name)
        for i, r in enumerate(rows, 1):
            v = _clean(r[0])
            if not v:
                continue
            # Plant og EQ_Cat har en ekstra kolonne (nummer / beskrivelse).
            # Den er det, brugeren ser; noeglen er det, der gemmes.
            extra = _clean(r[1]) if len(r) > 1 else ""
            values.append([list_name, v, extra, i])

    # SCEq foldes ud pr. klasse
    head, rows = table_rows(wb, "List data", "SCEq")
    for col, cls in enumerate(head):
        cls = _clean(cls)
        if not cls or cls == "SCEq":
            continue
        n = 0
        for r in rows:
            v = _clean(r[col]) if col < len(r) else ""
            if not v:
                continue
            n += 1
            values.append(["SCEq:" + cls, v, "", n])

    write_seed(os.path.join(SEED, "MD_FLValueList.csv"),
               ["ListName", "Value", "Value2", "Sort"], values)

    # ---- opslagstabeller -------------------------------------------------
    keys, dropped = [], 0
    for xl_name, (kind, ki, vi) in sorted(KEY_TABLES.items()):
        head, rows = table_rows(wb, "DictionaryTable", xl_name)
        # VBA laeser tabellen med dict(key) = value - en TILDELING, ikke
        # .Add. Dubletter overskriver derfor hinanden, og den sidste
        # vinder. Her gores det samme, saa listen indeholder praecis det
        # arket selv ender med.
        seen = {}
        order = []
        for r in rows:
            k = _clean(r[ki]).upper()
            if not k:
                continue
            val = _clean(r[vi]) if vi is not None and vi < len(r) else ""
            # BR18: noeglen er key12, vaerdien er den tilladte key17.
            # Beskrivelsen staar i tredje kolonne.
            desc = _clean(r[2]) if len(r) > 2 else ""
            # BR18 har flere raekker pr. key12 med vilje - det er en
            # vaerdiliste, ikke et opslag. Der maa ikke dedupliceres.
            uniq = (k, val) if kind == "BR18" else k
            if uniq in seen:
                dropped += 1
            else:
                order.append(uniq)
            seen[uniq] = [kind, k, val, desc]
        keys.extend(seen[u] for u in order)

    write_seed(os.path.join(SEED, "MD_FLKey.csv"),
               ["KeyType", "KeyValue", "Value", "Description"], keys)

    # ---- opsummering -----------------------------------------------------
    names = {}
    for r in values:
        names[r[0]] = names.get(r[0], 0) + 1
    print("MD_FLValueList.csv %4d vaerdier i %d lister" % (len(values), len(names)))
    for n in sorted(names):
        print("    %-28s %4d" % (n, names[n]))
    kinds = {}
    for r in keys:
        kinds[r[0]] = kinds.get(r[0], 0) + 1
    print("MD_FLKey.csv       %4d opslag (%d dubletter overskrevet)"
          % (len(keys), dropped))
    for n in sorted(kinds):
        print("    %-28s %4d" % (n, kinds[n]))


if __name__ == "__main__":
    main()
